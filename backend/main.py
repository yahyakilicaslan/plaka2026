"""
===================================================================
EVO SMART LPR BACKEND - v77.50 ULTRA-VISION
GÜNCELLEME: Telegram üzerinden butonlu kapı açma ve 
"/sorgula PLAKA" komutu ile detaylı plaka sorgulama özelliği eklendi.
===================================================================
"""
from fastapi import FastAPI, WebSocket, WebSocketDisconnect, HTTPException, UploadFile, File, BackgroundTasks
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from datetime import datetime, timedelta
from typing import Optional, List
from contextlib import asynccontextmanager
import json, uvicorn, os, sqlite3, requests, psutil, asyncio, csv, io, shutil, time
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMAGE_DIR = os.path.join(BASE_DIR, "captured_images")
FRONTEND_DIR = os.path.join(BASE_DIR, "frontend")
DB_PATH = os.path.join(BASE_DIR, "lpr.db")
DATA_DIR = os.path.join(BASE_DIR, "data")

if not os.path.exists(DATA_DIR): os.makedirs(DATA_DIR)
if not os.path.exists(IMAGE_DIR): os.makedirs(IMAGE_DIR)

async def cleanup_old_images():
    while True:
        try:
            conn = sqlite3.connect(DB_PATH)
            c = conn.cursor()
            c.execute("SELECT value FROM settings WHERE key='auto_delete_days'")
            res = c.fetchone()
            days = int(res[0]) if res else 7
            conn.close()
            if days > 0:
                now = time.time()
                cutoff = now - (days * 86400)
                for f in os.listdir(IMAGE_DIR):
                    f_path = os.path.join(IMAGE_DIR, f)
                    if os.path.isfile(f_path) and os.path.getmtime(f_path) < cutoff:
                        try: os.remove(f_path)
                        except: pass
        except Exception: pass
        await asyncio.sleep(86400)

# --- TELEGRAM ARKA PLAN DİNLEME (POLLING) MOTORU ---
async def telegram_polling_loop():
    last_update_id = 0
    while True:
        try:
            conn = sqlite3.connect(DB_PATH)
            c = conn.cursor()
            c.execute("SELECT value FROM settings WHERE key='telegram_token'")
            res = c.fetchone()
            token = res[0] if res else ""
            conn.close()
            
            if not token:
                await asyncio.sleep(10)
                continue

            url = f"https://api.telegram.org/bot{token}/getUpdates?offset={last_update_id}&timeout=10"
            loop = asyncio.get_event_loop()
            resp = await loop.run_in_executor(None, lambda: requests.get(url, timeout=15))
            
            if resp.status_code == 200:
                data = resp.json()
                for result in data.get("result", []):
                    last_update_id = result["update_id"] + 1
                    
                    # 1. BUTON TIKLAMALARI (Callback Query)
                    if "callback_query" in result:
                        cb = result["callback_query"]
                        cb_data = cb.get("data", "")
                        cb_id = cb.get("id")
                        chat_id = cb["message"]["chat"]["id"]
                        msg_id = cb["message"]["message_id"]

                        if cb_data.startswith("open_"):
                            slot = cb_data.split("_")[1]
                            conn = sqlite3.connect(DB_PATH)
                            c = conn.cursor()
                            g_info = c.execute("SELECT g.ip_address, g.endpoint FROM cameras c JOIN gates g ON c.gate_id = g.id WHERE c.slot = ? LIMIT 1", (slot,)).fetchone()
                            conn.close()
                            
                            if g_info:
                                try: await loop.run_in_executor(None, lambda: requests.get(f"http://{g_info[0]}{g_info[1]}", timeout=3))
                                except: pass
                            
                            new_cap = cb["message"].get("caption", "") + "\n\n✅ KAPI UZAKTAN AÇILDI."
                            await loop.run_in_executor(None, lambda: requests.post(f"https://api.telegram.org/bot{token}/editMessageCaption", json={"chat_id": chat_id, "message_id": msg_id, "caption": new_cap, "reply_markup": {"inline_keyboard": []}}))
                            await loop.run_in_executor(None, lambda: requests.post(f"https://api.telegram.org/bot{token}/answerCallbackQuery", json={"callback_query_id": cb_id, "text": "Kapı Başarıyla Açıldı!"}))
                        
                        elif cb_data.startswith("reject_"):
                            new_cap = cb["message"].get("caption", "") + "\n\n❌ GİRİŞ REDDEDİLDİ."
                            await loop.run_in_executor(None, lambda: requests.post(f"https://api.telegram.org/bot{token}/editMessageCaption", json={"chat_id": chat_id, "message_id": msg_id, "caption": new_cap, "reply_markup": {"inline_keyboard": []}}))
                            await loop.run_in_executor(None, lambda: requests.post(f"https://api.telegram.org/bot{token}/answerCallbackQuery", json={"callback_query_id": cb_id, "text": "Araç Reddedildi."}))
                    
                    # 2. METİN MESAJLARI VE SORGULAMA (Yeni Eklendi)
                    elif "message" in result and "text" in result["message"]:
                        text = result["message"]["text"].strip()
                        chat_id = result["message"]["chat"]["id"]
                        
                        if text.lower().startswith("/sorgula"):
                            parts = text.split(maxsplit=1)
                            if len(parts) > 1:
                                raw_plate = parts[1].upper().replace(" ", "")
                                
                                conn = sqlite3.connect(DB_PATH)
                                conn.row_factory = sqlite3.Row
                                c = conn.cursor()
                                
                                # Sistemde Kayıtlı Mı?
                                c.execute('''SELECT p.status, r.name, r.phone, b.name as block_name, s.name as site_name 
                                             FROM plates p 
                                             JOIN residents r ON p.resident_id = r.id 
                                             JOIN blocks b ON r.block_id = b.id 
                                             JOIN sites s ON b.site_id = s.id
                                             WHERE p.plate = ?''', (raw_plate,))
                                res_db = c.fetchone()
                                
                                # En Son Ne Zaman Geçti?
                                c.execute("SELECT timestamp, direction, camera_id FROM logs WHERE plate = ? ORDER BY id DESC LIMIT 1", (raw_plate,))
                                last_log = c.fetchone()
                                conn.close()
                                
                                msg_reply = f"🔎 *Plaka Sorgusu:* {raw_plate}\n\n"
                                if res_db:
                                    msg_reply += f"👤 *Sahibi:* {res_db['name']}\n"
                                    msg_reply += f"🏢 *Adres:* {res_db['site_name']} - {res_db['block_name']}\n"
                                    msg_reply += f"📱 *Tel:* {res_db['phone'] or '-'}\n"
                                    msg_reply += f"🔰 *Durum:* {res_db['status']}\n"
                                else:
                                    msg_reply += "⚠️ *Bu plaka sistemde kayıtlı değil (Misafir).* \n"
                                    
                                if last_log:
                                    msg_reply += f"\n⏱ *Son Görülme:* {last_log['timestamp']}\n"
                                    msg_reply += f"📍 *Kamera:* {last_log['camera_id']} ({last_log['direction']})"
                                else:
                                    msg_reply += "\n📉 *Sistemde hiç geçiş kaydı bulunamadı.*"
                                
                                await loop.run_in_executor(None, lambda: requests.post(f"https://api.telegram.org/bot{token}/sendMessage", json={"chat_id": chat_id, "text": msg_reply, "parse_mode": "Markdown"}, timeout=5))
                            else:
                                msg_reply = "Lütfen sorgulamak istediğiniz plakayı komutun yanına yazın. Örnek: `/sorgula 06DBN786`"
                                await loop.run_in_executor(None, lambda: requests.post(f"https://api.telegram.org/bot{token}/sendMessage", json={"chat_id": chat_id, "text": msg_reply, "parse_mode": "Markdown"}, timeout=5))

        except Exception:
            pass
        await asyncio.sleep(2)

@asynccontextmanager
async def lifespan(app: FastAPI):
    task = asyncio.create_task(broadcast_system_stats())
    cleanup_task = asyncio.create_task(cleanup_old_images())
    tel_task = asyncio.create_task(telegram_polling_loop())
    yield
    task.cancel(); cleanup_task.cancel(); tel_task.cancel()

app = FastAPI(lifespan=lifespan)
app.mount("/images", StaticFiles(directory=IMAGE_DIR), name="images")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_credentials=True, allow_methods=["*"], allow_headers=["*"])

class SiteModel(BaseModel): name: str
class BlockModel(BaseModel): site_id: int; name: str; flat_count: int
class ResidentModel(BaseModel): block_id: int; flat_number: int; name: str; phone: str; plates: List[dict]
class GateModel(BaseModel): name: str; ip_address: str; endpoint: str
class CameraModel(BaseModel): name: str; type: str; source: str; rtsp_user: Optional[str]=""; rtsp_pass: Optional[str]=""; slot: int; gate_id: Optional[int]=None; direction: Optional[str]="GIRIS"
class PlateLog(BaseModel): plate: str; confidence: float; camera_id: str; image_path: Optional[str]=None
class SettingsModel(BaseModel): 
    mode: str; rtsp_transport: Optional[str] = 'tcp'; frame_skip: Optional[int] = 0; ai_interval: Optional[float] = 0.04
    auto_delete_days: Optional[int] = 7; ocr_conf_threshold: Optional[float] = 0.40; vote_count: Optional[int] = 3
    vote_accept: Optional[int] = 2; debounce_time: Optional[int] = 15; api_worker_count: Optional[int] = 3
    tracking_enabled: Optional[bool] = False; gpu_acceleration: Optional[bool] = False; motion_sleep: Optional[bool] = False
    telegram_token: Optional[str] = ""; telegram_chat_id: Optional[str] = ""
    anti_blur: Optional[bool] = False; high_speed_tracker: Optional[bool] = False; dynamic_turbo: Optional[bool] = False

class QuotaSettingsModel(BaseModel): quota_enabled: bool; quota_limit: int; gate_trigger_enabled: bool; tabela_display_time: Optional[int] = 10; tabela_refresh_rate: Optional[int] = 500
class ROIModel(BaseModel): slot: int; roi: Optional[dict]; enabled: bool
class CamSettingsModel(BaseModel): slot: int; brightness: int = 0; contrast: float = 1.0

LAST_PLATE_INFO = {"plate": "", "owner": "", "status": "", "site": "", "timestamp": ""}

def init_db():
    conn = sqlite3.connect(DB_PATH); c = conn.cursor()
    try:
        c.execute("PRAGMA journal_mode=WAL")
        c.execute("PRAGMA synchronous=NORMAL")
        c.execute("PRAGMA temp_store=MEMORY")
        c.execute("PRAGMA cache_size=-20000")
    except Exception: pass
    c.execute('CREATE TABLE IF NOT EXISTS sites (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS blocks (id INTEGER PRIMARY KEY AUTOINCREMENT, site_id INTEGER, name TEXT, flat_count INTEGER)')
    c.execute('CREATE TABLE IF NOT EXISTS residents (id INTEGER PRIMARY KEY AUTOINCREMENT, block_id INTEGER, flat_number INTEGER, name TEXT, phone TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS plates (id INTEGER PRIMARY KEY AUTOINCREMENT, resident_id INTEGER, plate TEXT, status TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS gates (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, ip_address TEXT, endpoint TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS cameras (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT, type TEXT, source TEXT, rtsp_user TEXT, rtsp_pass TEXT, slot INTEGER, gate_id INTEGER, direction TEXT DEFAULT "GIRIS")')
    c.execute('CREATE TABLE IF NOT EXISTS logs (id INTEGER PRIMARY KEY AUTOINCREMENT, plate TEXT, timestamp TEXT, camera_id TEXT, image_path TEXT, status TEXT, owner TEXT, direction TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS parking_status (id INTEGER PRIMARY KEY AUTOINCREMENT, resident_id INTEGER, plate TEXT, is_inside INTEGER DEFAULT 0, entry_time TEXT)')
    c.execute('CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_plates_plate ON plates(plate)')
    c.execute('CREATE INDEX IF NOT EXISTS idx_logs_timestamp ON logs(timestamp)')
    
    defaults = [
        ('performance_mode', 'low'), ('quota_enabled', 'false'), ('quota_limit', '2'), ('gate_trigger_enabled', 'true'), ('rtsp_transport', 'tcp'), ('frame_skip', '2'),
        ('ai_interval', '0.06'), ('auto_delete_days', '7'), ('ocr_conf_threshold', '0.40'), ('vote_count', '3'), ('vote_accept', '2'), ('debounce_time', '15'),
        ('api_worker_count', '3'), ('tracking_enabled', 'true'), ('gpu_acceleration', 'false'), ('motion_sleep', 'false'), ('anti_blur', 'false'),
        ('high_speed_tracker', 'false'), ('dynamic_turbo', 'false')
    ]
    for k, v in defaults: c.execute("INSERT OR IGNORE INTO settings (key, value) VALUES (?, ?)", (k, v))
    conn.commit(); conn.close()

init_db()

class ConnectionManager:
    def __init__(self): self.active_connections: list[WebSocket] = []
    async def connect(self, websocket: WebSocket): await websocket.accept(); self.active_connections.append(websocket)
    def disconnect(self, websocket: WebSocket): self.active_connections.remove(websocket)
    async def broadcast(self, message: str):
        for c in self.active_connections:
            try: await c.send_text(message)
            except: pass
manager = ConnectionManager()

async def broadcast_system_stats():
    while True:
        try:
            msg = json.dumps({"type": "sys_stats", "cpu": psutil.cpu_percent(), "ram": psutil.virtual_memory().percent})
            await manager.broadcast(msg)
        except: pass
        await asyncio.sleep(5)

def _fire_nodemcu(url: str):
    """NodeMCU tetik (non-blocking arka plan). Kısa timeout."""
    try:
        requests.get(url, timeout=2)
    except Exception:
        pass

def send_telegram_alert_sync(token, chat_id, plate, status, owner, image_path, timestamp, cam_slot):
    try:
        url = f"https://api.telegram.org/bot{token}/sendPhoto"
        msg = f"🚨 DİKKAT! ONAY BEKLEYEN ARAÇ!\n\n🚗 Plaka: {plate}\n👤 Bilgi: {owner}\n⏱️ Tarih: {timestamp}\nSistem Durumu: {status}"
        
        reply_markup = json.dumps({
            "inline_keyboard": [
                [
                    {"text": "🟢 Kapıyı Aç", "callback_data": f"open_{cam_slot}"},
                    {"text": "🔴 Reddet", "callback_data": f"reject_{cam_slot}"}
                ]
            ]
        })

        real_img_path = os.path.join(IMAGE_DIR, os.path.basename(image_path))
        if os.path.exists(real_img_path):
            with open(real_img_path, 'rb') as f: 
                requests.post(url, data={'chat_id': chat_id, 'caption': msg, 'reply_markup': reply_markup}, files={'photo': f}, timeout=5)
        else: 
            requests.post(f"https://api.telegram.org/bot{token}/sendMessage", json={'chat_id': chat_id, 'text': msg, 'reply_markup': json.loads(reply_markup)}, timeout=5)
    except Exception: pass

@app.get("/api/settings")
def get_settings():
    conn = sqlite3.connect(DB_PATH); c = conn.cursor(); settings = {}
    for row in c.execute("SELECT key, value FROM settings").fetchall(): settings[row[0]] = row[1]
    conn.close()
    return {
        "mode": settings.get('performance_mode', 'low'), "quota_enabled": settings.get('quota_enabled', 'false') == 'true',
        "quota_limit": int(settings.get('quota_limit', '2')), "gate_trigger_enabled": settings.get('gate_trigger_enabled', 'true') == 'true',
        "tabela_display_time": int(settings.get('tabela_display_time', '10')), "tabela_refresh_rate": int(settings.get('tabela_refresh_rate', '500')),
        "rtsp_transport": settings.get('rtsp_transport', 'tcp'), "frame_skip": int(settings.get('frame_skip', '0')),
        "ai_interval": float(settings.get('ai_interval', '0.04')), "auto_delete_days": int(settings.get('auto_delete_days', '7')),
        "ocr_conf_threshold": float(settings.get('ocr_conf_threshold', '0.40')), "vote_count": int(settings.get('vote_count', '3')),
        "vote_accept": int(settings.get('vote_accept', '2')), "debounce_time": int(settings.get('debounce_time', '15')),
        "api_worker_count": int(settings.get('api_worker_count', '3')), "tracking_enabled": settings.get('tracking_enabled', 'false') == 'true',
        "gpu_acceleration": settings.get('gpu_acceleration', 'false') == 'true', "motion_sleep": settings.get('motion_sleep', 'false') == 'true',
        "telegram_token": settings.get('telegram_token', ''), "telegram_chat_id": settings.get('telegram_chat_id', ''),
        "anti_blur": settings.get('anti_blur', 'false') == 'true', "high_speed_tracker": settings.get('high_speed_tracker', 'false') == 'true',
        "dynamic_turbo": settings.get('dynamic_turbo', 'false') == 'true'
    }

@app.post("/api/settings")
def set_settings(s: SettingsModel):
    with sqlite3.connect(DB_PATH) as conn: 
        keys_vals = [
            ('performance_mode', s.mode), ('rtsp_transport', s.rtsp_transport), ('frame_skip', str(s.frame_skip)), 
            ('ai_interval', str(s.ai_interval)), ('auto_delete_days', str(s.auto_delete_days)), ('ocr_conf_threshold', str(s.ocr_conf_threshold)), 
            ('vote_count', str(s.vote_count)), ('vote_accept', str(s.vote_accept)), ('debounce_time', str(s.debounce_time)), 
            ('api_worker_count', str(s.api_worker_count)), ('tracking_enabled', str(s.tracking_enabled).lower()), 
            ('gpu_acceleration', str(s.gpu_acceleration).lower()), ('motion_sleep', str(s.motion_sleep).lower()), 
            ('telegram_token', s.telegram_token), ('telegram_chat_id', s.telegram_chat_id),
            ('anti_blur', str(s.anti_blur).lower()), ('high_speed_tracker', str(s.high_speed_tracker).lower()), ('dynamic_turbo', str(s.dynamic_turbo).lower())
        ]
        for k, v in keys_vals: conn.execute("REPLACE INTO settings (key, value) VALUES (?, ?)", (k, v))
    return {"status": "ok"}

@app.post("/api/settings/quota")
def set_quota_settings(s: QuotaSettingsModel):
    with sqlite3.connect(DB_PATH) as conn:
        for k, v in [('quota_enabled', str(s.quota_enabled).lower()), ('quota_limit', str(s.quota_limit)), ('gate_trigger_enabled', str(s.gate_trigger_enabled).lower()), ('tabela_display_time', str(s.tabela_display_time)), ('tabela_refresh_rate', str(s.tabela_refresh_rate))]:
            conn.execute("REPLACE INTO settings (key, value) VALUES (?, ?)", (k, v))
    return {"status": "ok"}

@app.post("/api/camera/roi")
def set_camera_roi(data: ROIModel):
    global CAMERA_ROI; CAMERA_ROI[data.slot] = {"roi": data.roi, "enabled": data.enabled}
    try:
        with open(os.path.join(DATA_DIR, "camera_roi.json"), "w") as f: json.dump(CAMERA_ROI, f)
    except: pass
    return {"status": "ok"}

@app.get("/api/camera/roi/{slot}")
def get_camera_roi(slot: int):
    global CAMERA_ROI
    try:
        with open(os.path.join(DATA_DIR, "camera_roi.json"), "r") as f: CAMERA_ROI = {int(k): v for k, v in json.load(f).items()}
    except: pass
    return CAMERA_ROI.get(slot, {"roi": None, "enabled": False})

@app.post("/api/camera/settings")
def set_camera_settings(data: CamSettingsModel):
    global CAMERA_SETTINGS; CAMERA_SETTINGS[data.slot] = {"brightness": data.brightness, "contrast": data.contrast}
    try:
        with open(os.path.join(DATA_DIR, "camera_settings.json"), "w") as f: json.dump(CAMERA_SETTINGS, f)
    except: pass
    return {"status": "ok"}

@app.get("/api/camera/settings/{slot}")
def get_camera_settings(slot: int):
    global CAMERA_SETTINGS
    try:
        with open(os.path.join(DATA_DIR, "camera_settings.json"), "r") as f: CAMERA_SETTINGS = {int(k): v for k, v in json.load(f).items()}
    except: pass
    return CAMERA_SETTINGS.get(slot, {"brightness": 0, "contrast": 1.0})

@app.get("/api/tabela/status")
def tabela_status(): return LAST_PLATE_INFO
@app.get("/api/led/status")
def led_status(): return LAST_PLATE_INFO

@app.get("/api/stats/today")
def get_today_stats():
    today = datetime.now().strftime("%d.%m.%Y")
    conn = sqlite3.connect(DB_PATH); c = conn.cursor()
    c.execute("SELECT status, count(*) FROM logs WHERE timestamp LIKE ? GROUP BY status", (today + "%",))
    data = c.fetchall(); conn.close()
    stats = {"total": 0, "allowed": 0, "guest": 0, "banned": 0}
    for st, count in data:
        stats["total"] += count
        if st == "IZINLI": stats["allowed"] += count
        elif st in ["YASAKLI", "KOTA_DOLU"]: stats["banned"] += count
        else: stats["guest"] += count
    return stats

@app.get("/api/stats/hourly")
def get_hourly_stats():
    today = datetime.now().strftime("%d.%m.%Y")
    conn = sqlite3.connect(DB_PATH); c = conn.cursor()
    c.execute("SELECT timestamp FROM logs WHERE timestamp LIKE ?", (today + "%",))
    rows = c.fetchall(); conn.close()
    hours = [0] * 24
    for r in rows:
        try: hours[int(r[0].split(" ")[1].split(":")[0])] += 1
        except: pass
    return hours

@app.post("/api/log-plate")
async def log_plate(data: PlateLog, background_tasks: BackgroundTasks):
    timestamp = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    c.execute("SELECT key, value FROM settings"); settings = {r[0]: r[1] for r in c.fetchall()}
    quota_enabled = settings.get('quota_enabled', 'false') == 'true'
    quota_limit = int(settings.get('quota_limit', '2'))
    gate_trigger_enabled = settings.get('gate_trigger_enabled', 'true') == 'true'
    tel_token = settings.get('telegram_token', ''); tel_chat = settings.get('telegram_chat_id', '')
    
    cam_direction = "GIRIS"
    slot_num = 1
    try:
        slot_num = int(data.camera_id.split("-")[1])
        cam_row = c.execute("SELECT direction FROM cameras WHERE slot = ?", (slot_num,)).fetchone()
        if cam_row: cam_direction = cam_row['direction'] or 'GIRIS'
    except: pass
    
    c.execute("SELECT p.status, p.resident_id, r.name as owner, b.name as block_name, s.name as site_name FROM plates p JOIN residents r ON p.resident_id = r.id JOIN blocks b ON r.block_id = b.id JOIN sites s ON b.site_id = s.id WHERE p.plate = ?", (data.plate,))
    res = c.fetchone()
    status = "MİSAFİR"; owner = "Tanımsız Araç"; site = ""; granted = False
    if res:
        status = res["status"]; owner = f"{res['block_name']} / {res['owner']}"; site = res["site_name"]
        if status == "IZINLI":
            if quota_enabled and cam_direction == "GIRIS":
                inside_count = c.execute("SELECT COUNT(*) as cnt FROM parking_status ps JOIN plates pl ON ps.plate = pl.plate WHERE pl.resident_id = ? AND ps.is_inside = 1", (res["resident_id"],)).fetchone()['cnt']
                if inside_count >= quota_limit: status = "KOTA_DOLU"
                else: granted = True
            else: granted = True
            if quota_enabled and granted and cam_direction == "GIRIS": c.execute("INSERT OR REPLACE INTO parking_status (resident_id, plate, is_inside, entry_time) VALUES (?, ?, 1, ?)", (res["resident_id"], data.plate, timestamp))
            elif quota_enabled and granted and cam_direction == "CIKIS": c.execute("UPDATE parking_status SET is_inside = 0 WHERE plate = ?", (data.plate,))
    
    c.execute("INSERT INTO logs (plate, timestamp, camera_id, image_path, status, owner, direction) VALUES (?, ?, ?, ?, ?, ?, ?)", (data.plate, timestamp, data.camera_id, data.image_path, status, owner, cam_direction))
    conn.commit()
    
    nodemcu_ip = ""; door_name = data.camera_id; nodemcu_url = ""
    try:
        gate_info = c.execute("SELECT g.ip_address, g.name as door_name, g.endpoint FROM cameras c JOIN gates g ON c.gate_id = g.id WHERE c.slot = ? LIMIT 1", (slot_num,)).fetchone()
        if gate_info:
            nodemcu_ip = gate_info['ip_address']; door_name = gate_info['door_name']
            nodemcu_url = f"http://{nodemcu_ip}{gate_info['endpoint']}"
            if granted and gate_trigger_enabled:
                # Kapı tetiğini arka plana at → yanıt anında dönsün
                background_tasks.add_task(_fire_nodemcu, nodemcu_url)
    except Exception: pass
    conn.close()
    
    global LAST_PLATE_INFO
    LAST_PLATE_INFO = {"plate": data.plate, "owner": owner, "status": status, "site": site, "timestamp": timestamp, "direction": cam_direction, "quota_message": "Otopark kotası dolu" if status == "KOTA_DOLU" else ""}
    
    if status in ['YASAKLI', 'KOTA_DOLU', 'MİSAFİR'] and tel_token and tel_chat: 
        background_tasks.add_task(send_telegram_alert_sync, tel_token, tel_chat, data.plate, status, owner, data.image_path, timestamp, slot_num)
        
    msg = json.dumps({"type": "log", "plate": data.plate, "time": timestamp, "camera": data.camera_id, "image_path": data.image_path, "status": status, "owner": owner, "gate_opened": granted, "direction": cam_direction})
    await manager.broadcast(msg); return {"status": "ok", "plate_status": status, "gate_opened": granted, "nodemcu_ip": nodemcu_ip}

@app.get("/api/reports")
def get_reports(start_date: str = None, end_date: str = None, plate: str = None, status: str = None):
    query = "SELECT * FROM logs WHERE 1=1"
    params = []
    if plate: query += " AND plate LIKE ?"; params.append(f"%{plate}%")
    if status: query += " AND status = ?"; params.append(status)
    query += " ORDER BY id DESC LIMIT 500"
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    c.execute(query, tuple(params)); res = [dict(row) for row in c.fetchall()]; conn.close()
    if start_date or end_date:
        filtered = []
        sd = datetime.strptime(start_date, "%Y-%m-%d").date() if start_date else None
        ed = datetime.strptime(end_date, "%Y-%m-%d").date() if end_date else None
        for r in res:
            try:
                rd = datetime.strptime(r['timestamp'].split(' ')[0], "%d.%m.%Y").date()
                if sd and rd < sd: continue
                if ed and rd > ed: continue
                filtered.append(r)
            except: filtered.append(r)
        return filtered
    return res

@app.get("/api/backup/download")
def download_backup():
    if os.path.exists(DB_PATH): return FileResponse(DB_PATH, filename=f"lpr_backup_{datetime.now().strftime('%Y%m%d')}.db", media_type='application/x-sqlite3')
    return {"error": "Veritabanı bulunamadı"}

@app.post("/api/backup/restore")
async def restore_backup(file: UploadFile = File(...)):
    try:
        if os.path.exists(DB_PATH): shutil.copy(DB_PATH, DB_PATH + ".bak")
        with open(DB_PATH, "wb") as buffer: shutil.copyfileobj(file.file, buffer)
        return {"status": "ok", "message": "Yedek başarıyla yüklendi. Lütfen sistemi yeniden başlatın."}
    except Exception as e: return {"status": "error", "message": str(e)}

@app.get("/api/export/template")
def get_template():
    wb = Workbook(); ws = wb.active; ws.title = "Sakinler"
    header_font = Font(bold=True, color="FFFFFF", size=11); header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    headers = ['Site Adı', 'Blok Adı', 'Daire No', 'Ad Soyad', 'Telefon', 'Plaka1', 'Durum1', 'Plaka2', 'Durum2', 'Plaka3', 'Durum3', 'Plaka4', 'Durum4', 'Plaka5', 'Durum5']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = thin_border
    for i, width in enumerate([15, 12, 10, 20, 15, 12, 8, 12, 8, 12, 8, 12, 8, 12, 8], 1): ws.column_dimensions[get_column_letter(i)].width = width
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={"Content-Disposition": "attachment; filename=sakinler_sablon.xlsx"})

@app.get("/api/export/residents")
def export_residents():
    try:
        conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
        c.execute("SELECT r.id, r.flat_number, r.name, r.phone, b.name as block_name, s.name as site_name FROM residents r JOIN blocks b ON r.block_id = b.id JOIN sites s ON b.site_id = s.id ORDER BY s.name, b.name, r.flat_number")
        residents = c.fetchall()
        wb = Workbook(); ws = wb.active; ws.title = "Sakinler"
        header_font = Font(bold=True, color="FFFFFF", size=11); header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
        headers = ['Site Adı', 'Blok Adı', 'Daire No', 'Ad Soyad', 'Telefon', 'Plaka1', 'Durum1', 'Plaka2', 'Durum2', 'Plaka3', 'Durum3', 'Plaka4', 'Durum4', 'Plaka5', 'Durum5']
        for col, header in enumerate(headers, 1): cell = ws.cell(row=1, column=col, value=header); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = thin_border
        for row_idx, res in enumerate(residents, 2):
            c.execute("SELECT plate, status FROM plates WHERE resident_id=?", (res['id'],)); plates = c.fetchall()
            row_data = [res['site_name'], res['block_name'], res['flat_number'], res['name'], res['phone'] or '']
            for i in range(5):
                if i < len(plates): row_data.extend([plates[i]['plate'], 1 if plates[i]['status'] == 'IZINLI' else 0])
                else: row_data.extend(['', ''])
            for col_idx, value in enumerate(row_data, 1): ws.cell(row=row_idx, column=col_idx, value=value).border = thin_border
        conn.close()
        output = io.BytesIO(); wb.save(output); output.seek(0)
        return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={"Content-Disposition": f"attachment; filename=sakinler_export_{datetime.now().strftime('%Y%m%d')}.xlsx"})
    except Exception as e: return {"status": "error", "message": str(e)}

@app.post("/api/import/residents")
async def import_residents(file: UploadFile = File(...)):
    try:
        content = await file.read(); filename = file.filename.lower(); conn = sqlite3.connect(DB_PATH); c = conn.cursor(); count = 0; errors = []
        if filename.endswith(('.xlsx', '.xls')):
            wb = load_workbook(io.BytesIO(content), data_only=True); ws = wb.active
            for line_num, row in enumerate(list(ws.iter_rows(min_row=2, values_only=True)), start=2):
                try:
                    if not row or len(row) < 5: continue
                    site_name, block_name, flat_num, name, phone = [str(row[i]).strip() if row[i] else '' for i in range(5)]
                    if not site_name or not block_name or not name or site_name == 'None': continue
                    plates_data = []
                    i = 5
                    while i < len(row) - 1:
                        plate_val, status_val = (row[i] if i < len(row) else None), (row[i+1] if i+1 < len(row) else None)
                        plate = str(plate_val).strip().upper().replace(" ", "") if plate_val and str(plate_val) != 'None' else ''
                        if plate:
                            status = 'IZINLI'
                            if status_val is not None and str(status_val) in ['0', '0.0', 'False', 'YASAKLI']: status = 'YASAKLI'
                            plates_data.append({'plate': plate, 'status': status})
                        i += 2
                    site_id = c.execute("SELECT id FROM sites WHERE name=?", (site_name,)).fetchone(); site_id = site_id[0] if site_id else c.execute("INSERT INTO sites (name) VALUES (?)", (site_name,)).lastrowid
                    block_id = c.execute("SELECT id FROM blocks WHERE site_id=? AND name=?", (site_id, block_name)).fetchone(); block_id = block_id[0] if block_id else c.execute("INSERT INTO blocks (site_id, name, flat_count) VALUES (?, ?, ?)", (site_id, block_name, 50)).lastrowid
                    res_id = c.execute("SELECT id FROM residents WHERE block_id=? AND flat_number=?", (block_id, flat_num)).fetchone()
                    if res_id: c.execute("UPDATE residents SET name=?, phone=? WHERE id=?", (name, phone, res_id[0])); res_id = res_id[0]
                    else: res_id = c.execute("INSERT INTO residents (block_id, flat_number, name, phone) VALUES (?, ?, ?, ?)", (block_id, flat_num, name, phone)).lastrowid
                    c.execute("DELETE FROM plates WHERE resident_id=?", (res_id,))
                    for p_data in plates_data[:5]: c.execute("INSERT INTO plates (resident_id, plate, status) VALUES (?, ?, ?)", (res_id, p_data['plate'], p_data['status']))
                    count += 1
                except Exception as row_error: errors.append(f"Satır {line_num}: {str(row_error)}")
        conn.commit(); conn.close()
        return {"status": "ok", "message": f"✅ {count} kayıt başarıyla aktarıldı."}
    except Exception as e: return {"status": "error", "message": f"Hata: {str(e)}"}

@app.get("/api/sites")
def get_sites():
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    c.execute("SELECT * FROM sites"); sites = [dict(row) for row in c.fetchall()]
    for s in sites: c.execute("SELECT * FROM blocks WHERE site_id = ?", (s["id"],)); s["blocks"] = [dict(row) for row in c.fetchall()]
    conn.close(); return sites

@app.post("/api/sites")
def add_site(site: SiteModel):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("INSERT INTO sites (name) VALUES (?)", (site.name,))
    return {"status": "ok"}
@app.put("/api/sites/{id}")
def update_site(id: int, site: SiteModel):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("UPDATE sites SET name=? WHERE id=?", (site.name, id))
    return {"status": "updated"}
@app.delete("/api/sites/{id}")
def del_site(id: int):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("DELETE FROM sites WHERE id=?", (id,)); conn.execute("DELETE FROM blocks WHERE site_id=?", (id,))
    return {"status": "ok"}

@app.post("/api/blocks")
def add_block(block: BlockModel):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("INSERT INTO blocks (site_id, name, flat_count) VALUES (?, ?, ?)", (block.site_id, block.name, block.flat_count))
    return {"status": "ok"}
@app.put("/api/blocks/{id}")
def update_block(id: int, block: BlockModel):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("UPDATE blocks SET name=?, flat_count=? WHERE id=?", (block.name, block.flat_count, id))
    return {"status": "updated"}

@app.get("/api/residents")
def get_res():
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    c.execute("SELECT r.*, b.name as block_name, b.id as block_id_real, s.name as site_name FROM residents r JOIN blocks b ON r.block_id = b.id JOIN sites s ON b.site_id = s.id ORDER BY r.id DESC")
    res = [dict(row) for row in c.fetchall()]
    for r in res: c.execute("SELECT plate, status FROM plates WHERE resident_id=?", (r["id"],)); r["plates"] = [dict(row) for row in c.fetchall()]
    conn.close(); return res

@app.post("/api/residents")
def add_res(res: ResidentModel):
    conn = sqlite3.connect(DB_PATH); c = conn.cursor()
    c.execute("INSERT INTO residents (block_id, flat_number, name, phone) VALUES (?, ?, ?, ?)", (res.block_id, res.flat_number, res.name, res.phone))
    rid = c.lastrowid
    for p in res.plates[:5]: c.execute("INSERT INTO plates (resident_id, plate, status) VALUES (?, ?, ?)", (rid, p['plate'].upper().replace(" ", ""), p['status']))
    conn.commit(); conn.close(); return {"status": "ok"}
@app.put("/api/residents/{id}")
def upd_res(id: int, res: ResidentModel):
    conn = sqlite3.connect(DB_PATH); c = conn.cursor()
    c.execute("UPDATE residents SET block_id=?, flat_number=?, name=?, phone=? WHERE id=?", (res.block_id, res.flat_number, res.name, res.phone, id))
    c.execute("DELETE FROM plates WHERE resident_id=?", (id,))
    for p in res.plates[:5]: 
        if p['plate']: c.execute("INSERT INTO plates (resident_id, plate, status) VALUES (?, ?, ?)", (id, p['plate'].upper().replace(" ", ""), p['status']))
    conn.commit(); conn.close(); return {"status": "updated"}
@app.delete("/api/residents/{id}")
def del_res_api(id: int):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("DELETE FROM residents WHERE id=?", (id,)); conn.execute("DELETE FROM plates WHERE resident_id=?", (id,))
    return {"status": "ok"}

@app.get("/api/gates")
def get_gates():
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    c.execute("SELECT * FROM gates"); res = [dict(row) for row in c.fetchall()]; conn.close(); return res

@app.post("/api/gates")
def add_gate(g: GateModel):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("INSERT INTO gates (name, ip_address, endpoint) VALUES (?,?,?)", (g.name, g.ip_address, g.endpoint))
    return {"status": "ok"}
@app.put("/api/gates/{id}")
def update_gate(id: int, g: GateModel):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("UPDATE gates SET name=?, ip_address=?, endpoint=? WHERE id=?", (g.name, g.ip_address, g.endpoint, id))
    return {"status": "updated"}
@app.delete("/api/gates/{id}")
def del_gate(id: int):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("DELETE FROM gates WHERE id=?", (id,))
    return {"status": "ok"}

@app.post("/api/gate/test/{gate_id}")
def test_gate(gate_id: int, background_tasks: BackgroundTasks):
    """Kapı test açma - dashboard 'Cihazlar' sekmesinden tetiklenir."""
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    row = c.execute("SELECT ip_address, endpoint, name FROM gates WHERE id=?", (gate_id,)).fetchone()
    conn.close()
    if not row:
        return {"status": "error", "message": "Kapı bulunamadı"}
    url = f"http://{row['ip_address']}{row['endpoint']}"
    background_tasks.add_task(_fire_nodemcu, url)
    return {"status": "ok", "message": f"{row['name']} tetiklendi", "url": url}

@app.post("/api/gate/manual-open/{slot}")
async def manual_open_gate(slot: int, background_tasks: BackgroundTasks):
    """Desktop uygulamadan manuel kapı açma. Slot üzerinden ilgili NodeMCU'yu tetikler."""
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    row = c.execute("""SELECT g.ip_address, g.endpoint, g.name as door_name, c.name as cam_name
                       FROM cameras c LEFT JOIN gates g ON c.gate_id = g.id
                       WHERE c.slot = ? LIMIT 1""", (slot,)).fetchone()
    conn.close()
    if not row or not row['ip_address']:
        return {"status": "error", "message": f"Slot {slot} için tanımlı kapı yok"}
    url = f"http://{row['ip_address']}{row['endpoint']}"
    background_tasks.add_task(_fire_nodemcu, url)
    timestamp = datetime.now().strftime("%d.%m.%Y %H:%M:%S")
    # Dashboard'ı bilgilendir
    msg = json.dumps({"type": "manual_open", "slot": slot, "camera": row['cam_name'],
                      "door": row['door_name'], "time": timestamp})
    await manager.broadcast(msg)
    return {"status": "ok", "message": f"{row['door_name']} manuel açıldı", "url": url}

@app.get("/api/cameras/slot/{slot}")
def get_camera_by_slot(slot: int):
    """Desktop uygulama için hızlı kamera bilgisi (gate dahil)."""
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    row = c.execute("""SELECT c.*, g.name as gate_name, g.ip_address as gate_ip, g.endpoint as gate_endpoint
                       FROM cameras c LEFT JOIN gates g ON c.gate_id = g.id WHERE c.slot=?""", (slot,)).fetchone()
    conn.close()
    return dict(row) if row else {}

@app.get("/api/cameras")
def get_cams():
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    c.execute("SELECT c.*, g.name as gate_name FROM cameras c LEFT JOIN gates g ON c.gate_id=g.id ORDER BY c.slot ASC")
    res = [dict(row) for row in c.fetchall()]; conn.close(); return res

@app.post("/api/cameras")
def add_cam(cm: CameraModel):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("DELETE FROM cameras WHERE slot=?", (cm.slot,)) 
        conn.execute("INSERT INTO cameras (name, type, source, rtsp_user, rtsp_pass, slot, gate_id, direction) VALUES (?,?,?,?,?,?,?,?)", 
                    (cm.name, cm.type, cm.source, cm.rtsp_user, cm.rtsp_pass, cm.slot, cm.gate_id, cm.direction or 'GIRIS'))
    return {"status": "ok"}
@app.put("/api/cameras/{id}")
def update_cam(id: int, cm: CameraModel):
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute("UPDATE cameras SET name=?, type=?, source=?, rtsp_user=?, rtsp_pass=?, slot=?, gate_id=?, direction=? WHERE id=?", 
                    (cm.name, cm.type, cm.source, cm.rtsp_user, cm.rtsp_pass, cm.slot, cm.gate_id, cm.direction or 'GIRIS', id))
    return {"status": "updated"}
@app.delete("/api/cameras/{id}")
def del_cam(id: int):
    with sqlite3.connect(DB_PATH) as conn: conn.execute("DELETE FROM cameras WHERE id=?", (id,))
    return {"status": "ok"}

@app.get("/api/history")
async def get_history():
    conn = sqlite3.connect(DB_PATH); conn.row_factory = sqlite3.Row; c = conn.cursor()
    c.execute("SELECT * FROM logs ORDER BY id DESC LIMIT 50"); h = [dict(row) for row in c.fetchall()]; conn.close(); return h

@app.get("/")
async def get_dashboard():
    html_path = os.path.join(FRONTEND_DIR, "dashboard.html")
    if os.path.exists(html_path):
        with open(html_path, "r", encoding="utf-8") as f: return HTMLResponse(content=f.read())
    return HTMLResponse(content="Dashboard Yok!", status_code=404)

@app.get("/tabela")
async def get_tabela():
    html_path = os.path.join(FRONTEND_DIR, "tabela.html")
    if os.path.exists(html_path):
        with open(html_path, "r", encoding="utf-8") as f: return HTMLResponse(content=f.read())
    return HTMLResponse(content="Tabela sayfası bulunamadı!", status_code=404)

@app.websocket("/ws")
async def websocket_endpoint(websocket: WebSocket):
    await manager.connect(websocket)
    try: 
        while True: await websocket.receive_text()
    except WebSocketDisconnect: manager.disconnect(websocket)

if __name__ == "__main__":
    print("=" * 60)
    print("   EVO SMART LPR BACKEND - v77.50 ULTRA-VISION")
    print(f"   HTTP dinleyici: http://0.0.0.0:8000  (IPv4)")
    print(f"   Desktop/LPR Engine icin: http://127.0.0.1:8000")
    print("=" * 60)
    uvicorn.run(app, host="0.0.0.0", port=8000)